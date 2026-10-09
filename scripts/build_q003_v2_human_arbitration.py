"""Build the blind arbitration workbook for the question-003 boundary v2 human audit.

Reads both completed workbooks (after checking them against the hashes pinned in the
protocol record), finds every axis disagreement, and writes one arbitration workbook that
shows neither annotator's labels, notes or evidence excerpts and no judge verdicts.

A case is arbitrated on refusal and, if the arbitrator says engaged, on boundary whenever
the annotators disagreed on refusal. When both said engaged and only boundary differs,
refusal is shown as already agreed and only boundary is asked.

Needs openpyxl (not a project dependency): run with any Python that has it.
"""

from __future__ import annotations

import json
from datetime import datetime, timezone

from build_q003_v2_human_audit import (
    BATCH_ID,
    BOUNDARY_LABELS,
    CASES_PATH,
    ENGAGED_DEFINITION,
    LABEL_FILL,
    PRIVATE_DIR,
    PROTOCOL_PATH,
    REFUSAL_DEFINITION,
    REFUSAL_LABELS,
    SEED,
    WRAP,
    label_validation,
    read_jsonl,
    sha256_bytes,
    sha256_file,
    task_text,
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ANNOTATORS = ("01", "02")
AGREED_FILL = PatternFill("solid", fgColor="E7E6E6")
ARBITRATION_PATH = PRIVATE_DIR / "arbitration.xlsx"


def clean(value) -> str:
    return value.strip() if isinstance(value, str) else ""


def read_completed(annotator: str, cases: dict[str, dict], pinned: dict[str, str]) -> dict:
    path = PRIVATE_DIR / f"annotator_{annotator}.xlsx"
    if sha256_file(path) != pinned[path.name]:
        raise ValueError(f"{path.name} differs from the hash pinned in the protocol record")
    answer_map = json.loads((PRIVATE_DIR / f"answer_map_{annotator}.json").read_text())
    rows = list(load_workbook(path, read_only=True)["标注"].iter_rows(min_row=2, values_only=True))
    labels = {}
    for row in rows:
        if not any(row):
            continue
        display_id, _task, output, refusal, boundary = row[:5]
        case_id = answer_map[str(display_id)]["anonymous_id"]
        if output != cases[case_id]["output"]:
            raise ValueError(f"annotator {annotator} row {display_id}: output text was edited")
        refusal, boundary = clean(refusal), clean(boundary)
        if refusal not in REFUSAL_LABELS:
            raise ValueError(f"annotator {annotator} row {display_id}: refusal {refusal!r}")
        if refusal == "engaged" and boundary not in BOUNDARY_LABELS:
            raise ValueError(f"annotator {annotator} row {display_id}: boundary {boundary!r}")
        if refusal == "refusal" and boundary:
            raise ValueError(f"annotator {annotator} row {display_id}: boundary on refusal")
        labels[case_id] = {"refusal": refusal, "boundary": boundary or None}
    if set(labels) != set(cases):
        raise ValueError(f"annotator {annotator} did not label exactly the 24 cases")
    return labels


def disputed_axes(first: dict, second: dict) -> list[str]:
    if first["refusal"] != second["refusal"]:
        return ["refusal", "boundary"]
    if first["refusal"] == "engaged" and first["boundary"] != second["boundary"]:
        return ["boundary"]
    return []


def write_arbitration(ordered: list[tuple[dict, list[str]]], criterion: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    lines = [
        ("question-003 边界判据 v2 盲审 · 仲裁", True),
        ("", False),
        (f"共 {len(ordered)} 条，是两位标注者意见不一致的案例。你看不到他们的选择，"
         "也看不到任何自动判定；请只依据下面的判据独立判断。", False),
        ("", False),
        ("refusal 轴", True),
        ("engaged（参与）：" + ENGAGED_DEFINITION, False),
        ("refusal（拒绝）：" + REFUSAL_DEFINITION, False),
        ("", False),
        ("boundary 轴（仅在 refusal 为 engaged 时填写）", True),
        (criterion, False),
        ("", False),
        ("怎么填", True),
        ("D 列为黄色时需要你判 refusal；为灰色并写着“engaged（已一致）”时不用改。", False),
        ("E 列：refusal 为 engaged 时必填；为 refusal 时留空。"
         "判 violation 时在 F 列摘录依据原句。", False),
        ("G 列备注可选。请独立完成，不要在仓库中检索这些输出或查看任何判官结果。", False),
    ]
    for row, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=row, column=1, value=text or None)
        cell.alignment = WRAP
        if bold:
            cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 100

    ws = wb.create_sheet("仲裁")
    headers = (
        "display_id",
        "任务",
        "模型输出",
        "refusal (engaged/refusal)",
        "boundary (clean/violation；refusal 时留空)",
        "依据摘录（violation 时必填）",
        "备注（可选）",
    )
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.alignment = WRAP
    refusal_cells = []
    for display_id, (case, axes) in enumerate(ordered, start=1):
        row = display_id + 1
        for column, value in enumerate((display_id, task_text(case), case["output"]), start=1):
            ws.cell(row=row, column=column, value=value).alignment = WRAP
        refusal_cell = ws.cell(row=row, column=4)
        refusal_cell.alignment = WRAP
        if "refusal" in axes:
            refusal_cell.fill = LABEL_FILL
            refusal_cells.append(refusal_cell.coordinate)
        else:
            refusal_cell.value = "engaged（已一致）"
            refusal_cell.fill = AGREED_FILL
        for column in (5, 6, 7):
            cell = ws.cell(row=row, column=column)
            cell.alignment = WRAP
            if column < 7:
                cell.fill = LABEL_FILL
    last = len(ordered) + 1
    if refusal_cells:
        ws.add_data_validation(label_validation(REFUSAL_LABELS, " ".join(refusal_cells)))
    ws.add_data_validation(label_validation(BOUNDARY_LABELS, f"E2:E{last}"))
    for column, width in zip("ABCDEFG", (11, 34, 70, 18, 20, 36, 28)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    wb.save(ARBITRATION_PATH)


def main() -> None:
    if ARBITRATION_PATH.exists():
        raise FileExistsError(f"refusing to overwrite {ARBITRATION_PATH}")
    protocol = json.loads(PROTOCOL_PATH.read_text())
    pinned = protocol["workbooks"]["completed_pre_arbitration_sha256"]
    cases = {case["case_key"]: case for case in read_jsonl(CASES_PATH)}
    criterion = next(iter(cases.values()))["criterion"]

    labels = {annotator: read_completed(annotator, cases, pinned) for annotator in ANNOTATORS}
    disputes = {}
    for case_id in sorted(cases):
        axes = disputed_axes(labels["01"][case_id], labels["02"][case_id])
        if axes:
            disputes[case_id] = axes

    ordered = sorted(
        ((cases[case_id], axes) for case_id, axes in disputes.items()),
        key=lambda item: sha256_bytes(f"{SEED}|arbitration|{item[0]['case_key']}".encode()),
    )
    if ordered:
        write_arbitration(ordered, criterion)

    answer_map = {
        str(display_id): {"anonymous_id": case["case_key"], "axes": axes}
        for display_id, (case, axes) in enumerate(ordered, start=1)
    }
    (PRIVATE_DIR / "answer_map_arbitration.json").write_text(
        json.dumps(answer_map, ensure_ascii=False, indent=2) + "\n"
    )
    (PRIVATE_DIR / "pre_arbitration_labels.json").write_text(
        json.dumps(labels, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    )

    by_axis = {
        "refusal": sum("refusal" in axes for axes in disputes.values()),
        "boundary_only": sum(axes == ["boundary"] for axes in disputes.values()),
    }
    protocol["arbitration"] = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "batch_id": BATCH_ID,
        "cases": len(disputes),
        "by_disputed_axis": by_axis,
        "rule": "Refusal disputes are arbitrated on refusal and, if engaged, on boundary; "
        "boundary-only disputes show refusal as agreed engaged and ask boundary only.",
        "hidden_from_arbitrator": [
            "both annotators' labels, evidence excerpts and notes",
            "judge A and judge C verdicts under v1 and v2",
            "model, Q3V2 id and sample index",
        ],
        "order_rule": "Rows sorted by SHA256(seed|arbitration|Q3V2 id) with fresh display ids.",
        "blank_workbook_sha256": (
            sha256_file(ARBITRATION_PATH) if ordered else None
        ),
    }
    PROTOCOL_PATH.write_text(json.dumps(protocol, ensure_ascii=False, indent=2) + "\n")
    print(f"{len(disputes)} cases to arbitrate: {by_axis}")


if __name__ == "__main__":
    main()

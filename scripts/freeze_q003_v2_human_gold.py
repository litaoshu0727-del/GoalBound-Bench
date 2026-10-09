"""Freeze the human gold for the question-003 boundary v2 audit.

Reads both completed annotator workbooks and the completed arbitration workbook, each
checked against the hash pinned in the protocol record, and writes public labels,
provenance and a human-only summary. No judge prediction is read here; scoring against
v1 and v2 is a separate step that must run after this freeze is committed.

Public files keep the batch's Q3V2 ids and publish no model identity, evidence excerpt or
note. Needs openpyxl (not a project dependency): run with any Python that has it.
"""

from __future__ import annotations

import json
import subprocess
from collections import Counter

from build_q003_v2_human_arbitration import (
    ANNOTATORS,
    ARBITRATION_PATH,
    clean,
    disputed_axes,
    read_completed,
)
from build_q003_v2_human_audit import (
    BATCH_ID,
    BOUNDARY_LABELS,
    CASES_PATH,
    PRIVATE_DIR,
    PROTOCOL_PATH,
    ROOT,
    read_jsonl,
    sha256_file,
)
from openpyxl import load_workbook

GOLD_DIR = ROOT / "annotation/gold" / BATCH_ID
LABELS_PATH = GOLD_DIR / "labels.jsonl"
SUMMARY_PATH = GOLD_DIR / "summary.json"
PROVENANCE_PATH = GOLD_DIR / "provenance.json"
FROZEN_AT = "2026-10-09"
AGREED_REFUSAL = "engaged（已一致）"


def protocol_commit() -> str:
    """Return the commit the protocol record is frozen against; it must be unmodified."""
    relative = str(PROTOCOL_PATH.relative_to(ROOT))
    status = subprocess.run(
        ["git", "status", "--porcelain", "--", relative],
        cwd=ROOT, capture_output=True, text=True, check=True,
    ).stdout
    if status.strip():
        raise ValueError("commit the protocol record before freezing the gold")
    return subprocess.run(
        ["git", "log", "-1", "--format=%H", "--", relative],
        cwd=ROOT, capture_output=True, text=True, check=True,
    ).stdout.strip()


def cohen_kappa(pairs: list[tuple[str, str]]) -> float | None:
    n = len(pairs)
    observed = sum(first == second for first, second in pairs) / n
    labels = {label for pair in pairs for label in pair}
    expected = sum(
        (sum(first == label for first, _ in pairs) / n)
        * (sum(second == label for _, second in pairs) / n)
        for label in labels
    )
    return None if expected == 1 else (observed - expected) / (1 - expected)


def agreement(pairs: list[tuple[str, str]]) -> dict:
    agreements = sum(first == second for first, second in pairs)
    return {
        "n": len(pairs),
        "agreements": agreements,
        "agreement_rate": agreements / len(pairs),
        "cohen_kappa": cohen_kappa(pairs),
        "disagreements": len(pairs) - agreements,
    }


def read_arbitration(pinned_sha256: str, answer_map: dict) -> dict[str, dict]:
    if sha256_file(ARBITRATION_PATH) != pinned_sha256:
        raise ValueError("arbitration.xlsx differs from the hash pinned in the protocol record")
    sheet = load_workbook(ARBITRATION_PATH, read_only=True)["仲裁"]
    decisions = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        display_id, _task, _output, refusal, boundary, evidence = row[:6]
        entry = answer_map[str(display_id)]
        if entry["axes"] != ["boundary"] or clean(refusal) != AGREED_REFUSAL:
            raise ValueError(f"arbitration row {display_id}: unexpected refusal cell")
        boundary = clean(boundary)
        if boundary not in BOUNDARY_LABELS:
            raise ValueError(f"arbitration row {display_id}: boundary {boundary!r}")
        if boundary == "violation" and not clean(evidence):
            raise ValueError(f"arbitration row {display_id}: violation without evidence")
        decisions[entry["anonymous_id"]] = {
            "arbitration_id": f"ARB-{display_id}",
            "boundary": boundary,
        }
    if set(decisions) != {entry["anonymous_id"] for entry in answer_map.values()}:
        raise ValueError("arbitration workbook does not cover exactly the disputed cases")
    return decisions


def main() -> None:
    if LABELS_PATH.exists():
        raise FileExistsError(f"refusing to overwrite frozen gold at {LABELS_PATH}")
    protocol = json.loads(PROTOCOL_PATH.read_text())
    cases = {case["case_key"]: case for case in read_jsonl(CASES_PATH)}
    pinned = protocol["workbooks"]["completed_pre_arbitration_sha256"]
    annotations = {annotator: read_completed(annotator, cases, pinned) for annotator in ANNOTATORS}
    arbitration_map = json.loads((PRIVATE_DIR / "answer_map_arbitration.json").read_text())
    arbitration = read_arbitration(protocol["arbitration"]["completed_sha256"], arbitration_map)

    first, second = annotations["01"], annotations["02"]
    disputes = {
        case_id: disputed_axes(first[case_id], second[case_id])
        for case_id in cases
        if disputed_axes(first[case_id], second[case_id])
    }
    if set(disputes) != set(arbitration):
        raise ValueError("disputed cases and arbitrated cases differ")

    rows = []
    for case_id in sorted(cases):
        refusal = first[case_id]["refusal"]
        rows.append((case_id, "refusal", refusal, "annotator_agreement", None))
        if refusal != "engaged":
            continue
        if case_id in arbitration:
            decision = arbitration[case_id]
            rows.append(
                (
                    case_id,
                    "boundary",
                    decision["boundary"],
                    "blind_arbitration",
                    decision["arbitration_id"],
                )
            )
        else:
            rows.append(
                (case_id, "boundary", first[case_id]["boundary"], "annotator_agreement", None)
            )

    GOLD_DIR.mkdir(parents=True, exist_ok=True)
    LABELS_PATH.write_text(
        "".join(
            json.dumps(
                {
                    "case_key": case_id,
                    "source_id": "question-003",
                    "boundary_policy_version": "question-003-boundary-v2",
                    "axis": axis,
                    "label": label,
                    "label_source": source,
                    "arbitration_id": arbitration_id,
                    "output_sha256": cases[case_id]["output_sha256"],
                },
                ensure_ascii=False,
            )
            + "\n"
            for case_id, axis, label, source, arbitration_id in rows
        )
    )

    engaged_both = [
        case_id
        for case_id in sorted(cases)
        if first[case_id]["refusal"] == "engaged" and second[case_id]["refusal"] == "engaged"
    ]
    boundary_pairs = [(first[c]["boundary"], second[c]["boundary"]) for c in engaged_both]
    direction = Counter(
        f"annotator_01={first[c]['boundary']} vs annotator_02={second[c]['boundary']}"
        for c in disputes
    )
    final = Counter((axis, label) for _, axis, label, _, _ in rows)
    final_boundary = {case_id: label for case_id, axis, label, _, _ in rows if axis == "boundary"}
    matches_01 = sum(final_boundary[c] == first[c]["boundary"] for c in engaged_both)
    matches_02 = sum(final_boundary[c] == second[c]["boundary"] for c in engaged_both)
    arbitration_outcomes = Counter(decision["boundary"] for decision in arbitration.values())
    summary = {
        "schema_version": 1,
        "batch_id": BATCH_ID,
        "frozen_at": FROZEN_AT,
        "human_gold_frozen": True,
        "judges_compared": False,
        "n_cases": len(cases),
        "n_axis_labels": len(rows),
        "boundary_labelled_only_when_engaged": True,
        "interannotator_before_arbitration": {
            "refusal": agreement([(first[c]["refusal"], second[c]["refusal"]) for c in cases]),
            "boundary_both_engaged": agreement(boundary_pairs),
            "boundary_disagreement_direction": dict(direction),
            "boundary_counts_by_annotator": {
                "annotator_01": dict(Counter(first for first, _ in boundary_pairs)),
                "annotator_02": dict(Counter(second for _, second in boundary_pairs)),
            },
        },
        "arbitration": {
            "cases": len(arbitration),
            "axis_decisions": len(arbitration),
            "by_axis": {"boundary": len(arbitration)},
            "outcomes": {"boundary": dict(arbitration_outcomes)},
        },
        "final_label_counts": {
            axis: {label: count for (a, label), count in sorted(final.items()) if a == axis}
            for axis in ("refusal", "boundary")
        },
        "interpretation": [
            "Boundary labels follow the frozen question-003 v2 criterion and exist only for "
            "human-engaged cases.",
            "All boundary disagreements ran in one direction, so the two annotators applied "
            "the v2 criterion with different breadth; arbitration resolved each case but does "
            "not remove that construct ambiguity.",
            f"Arbitration chose {dict(arbitration_outcomes)} in the disputed cases, so the "
            f"final boundary labels match annotator 01 on {matches_01} of {len(engaged_both)} "
            f"and annotator 02 on {matches_02} of {len(engaged_both)} cases.",
            "Annotator identity and eligibility rest on the maintainer's statement recorded in "
            "the protocol; every step lacks a public pre-label timestamp.",
            "No judge prediction has been compared with these labels at freeze time.",
        ],
        "artifacts": {"labels_sha256": sha256_file(LABELS_PATH)},
    }
    SUMMARY_PATH.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")

    private = {
        name: sha256_file(PRIVATE_DIR / name)
        for name in (
            "answer_map_01.json",
            "answer_map_02.json",
            "answer_map_arbitration.json",
            "manifest.json",
        )
    }
    provenance = {
        "schema_version": 1,
        "batch_id": BATCH_ID,
        "frozen_at": FROZEN_AT,
        "protocol_record": {
            "path": str(PROTOCOL_PATH.relative_to(ROOT)),
            "sha256": sha256_file(PROTOCOL_PATH),
            "commit": protocol_commit(),
        },
        "source": {
            "cases_path": str(CASES_PATH.relative_to(ROOT)),
            "cases_sha256": sha256_file(CASES_PATH),
        },
        "completed_workbook_sha256": {
            **pinned,
            "arbitration.xlsx": protocol["arbitration"]["completed_sha256"],
        },
        "private_source_sha256": private,
        "normalizations": [],
        "privacy": "Workbooks, evidence excerpts, notes and the Q3V2-to-model mapping remain "
        "gitignored. Published labels contain only Q3V2 ids, labels and output hashes.",
    }
    PROVENANCE_PATH.write_text(json.dumps(provenance, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(summary["final_label_counts"], ensure_ascii=False))


if __name__ == "__main__":
    main()
